"""PPT Skill 接口与构建器（Phase X，§23）。

核心工作流只提供内容结构/章节/图表/结论/受众/页数/风格（PPTInput），
由本构建器或外部 PPT Skill 生成 slides.pptx。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

from ..errors import GenerationError
from ..logging_setup import get_logger
from ..workspace import WorkspaceManager

logger = get_logger("ppt")


@dataclass
class PPTSlide:
    """一页：目标 + 要点 + 可选真实图表。"""

    title: str
    bullets: list[str] = field(default_factory=list)
    figure: str = ""  # 相对 workspace 的图片路径
    notes: str = ""


@dataclass
class PPTInput:
    """X001–X006 的内容包。"""

    title: str
    slides: list[PPTSlide] = field(default_factory=list)
    audience: str = "答辩委员会"
    pages_hint: int = 10
    style: dict[str, Any] = field(default_factory=lambda: {"theme": "academic-blue"})


class PPTBuilder:
    """X007 调用点：默认内置 python-pptx 实现；可替换为外部 Skill。"""

    def __init__(self, ws: WorkspaceManager, external_skill=None):
        self.ws = ws
        self._external = external_skill  # Callable[[PPTInput, Path], Path]

    def build(self, ppt_input: PPTInput, out_rel: str = "slides/slides.pptx") -> Path:
        if self._external is not None:
            return self._external(ppt_input, self.ws.resolve(out_rel))
        try:
            from pptx import Presentation
            from pptx.util import Pt
        except ImportError as e:
            raise GenerationError("需要 python-pptx") from e
        prs = Presentation()
        # 封面
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = ppt_input.title
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = f"受众：{ppt_input.audience}"
        # 内容页
        for s in ppt_input.slides:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = s.title
            body = slide.placeholders[1].text_frame
            first = True
            for b in s.bullets:
                para = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                para.text = b
                para.font.size = Pt(18)
            if s.figure:
                img = self.ws.root / s.figure
                if img.is_file():
                    slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width / 3)
        out = self.ws.resolve(out_rel)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        return out

    # ---------- X008–X010 验证 ----------
    def verify(self, ppt_path: str | Path, ppt_input: PPTInput) -> dict:
        from pptx import Presentation

        p = Path(ppt_path)
        if not p.is_file() or p.stat().st_size == 0:
            return {"ok": False, "reason": "PPT 文件缺失或为空"}
        try:
            prs = Presentation(str(p))
        except Exception as e:
            return {"ok": False, "reason": f"无法解析: {e}"}
        n = len(prs.slides)
        texts = []
        for sl in prs.slides:
            for shape in sl.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        blob = "\n".join(texts)
        missing = [s.title for s in ppt_input.slides if s.title not in blob]
        ok = n >= len(ppt_input.slides) + 1 and not missing
        return {
            "ok": ok,
            "slides": n,
            "missing_sections": missing,
            "reason": "" if ok else f"页数或章节覆盖不足：缺 {missing}",
        }

    # ---------- 从 workspace 内容包生成输入（X002–X006） ----------
    @staticmethod
    def input_from_workspace_payload(payload: dict) -> PPTInput:
        slides = [
            PPTSlide(title=s["title"], bullets=s.get("bullets", []), figure=s.get("figure", ""))
            for s in payload.get("sections", [])
        ]
        return PPTInput(
            title=payload.get("title", "研究汇报"),
            slides=slides,
            audience=payload.get("audience", "答辩委员会"),
            pages_hint=payload.get("pages_hint", 10),
        )


def load_ppt_payload(ws: WorkspaceManager, rel: str = "slides/ppt-content.yaml") -> dict:
    p = ws.root / rel
    if not p.is_file():
        raise GenerationError(f"PPT 内容包不存在: {rel}")
    return yaml.safe_load(p.read_text(encoding="utf-8")) or {}
